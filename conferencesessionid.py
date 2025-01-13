import os

# Установка необходимых библиотек
os.system("pip install requests")
os.system("pip install openpyxl")
os.system("pip install selenium")
os.system("pip install webdriver-manager")

import openpyxl
import time
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from openpyxl import load_workbook
from webdriver_manager.chrome import ChromeDriverManager

input("Для запуска программы нажмите Enter")

# Путь к Вашему Excel-файлу
exl_path = r"ids.xlsx"
exl = load_workbook(exl_path)

# Работа с листом data
sheet_data = exl["data"]

login_exl = sheet_data.cell(row=2, column=1).value  # Логин
password_exl = sheet_data.cell(row=2, column=2).value  # Пароль

# Работа с листом mod_links
sheet_vk_mod = exl["mod_links"]

# Работа с листом ids_links
sheet_vk_ids = exl["ids_links"]

# Инициализация WebDriver
driver = webdriver.Chrome(service=Service(ChromeDriverManager().install()))

# Вход на сайт
driver.get("https://edu.vkurse.ru/v2/login")

# Кнопка Входа Synergy ID
login_button = WebDriverWait(driver, 10).until(
    EC.element_to_be_clickable((By.LINK_TEXT, "Войти через Synergy ID"))
)
login_button.click()

# Ввод логина
username_field = driver.find_element(By.ID, 'username')  # Updated to correct ID
username_field.send_keys(login_exl)

# Ввод пароля
password_field = driver.find_element(By.ID, 'password')  # Updated to correct ID
password_field.send_keys(password_exl)

# Кнопка Входа
login_button2 = WebDriverWait(driver, 10).until(
    EC.element_to_be_clickable((By.ID, 'kc-login'))  # Corrected to use EC.element_to_be_clickable
)
login_button2.click()

# Пропустить адрес после входа
driver.get("https://edu.vkurse.ru/v2/iva/home/chats?conferenceSessionOpenMode=PIP&currentChatId=21f6df6a-c56d-3d18-a630-e2afa279831f")

# Initialize the list to store second links
ids_links = []

# Перебор всех строк на листе mod_links
for row in sheet_vk_mod.iter_rows(min_row=2, values_only=True):
    mod_link = row[0]  # Assuming the mod_links are in the first column
    driver.get(mod_link)

    # Click the button with text "Просмотр материалов"
    material_button = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.XPATH, "//button[contains(@class, 'view-materials-button') and contains(., 'Просмотр материалов')]"))
    )
    material_button.click()  # Click the button

    # Wait for the page to load and get the second link from the address bar
    time.sleep(2)  # Adjust sleep time as necessary
    second_link = driver.current_url  # Get the current URL from the address bar
    ids_links.append(second_link)

# Save the second links back to the ids.xlsx file
for index, link in enumerate(ids_links, start=2):  # Start from row 2 to avoid overwriting headers
    sheet_vk_ids.cell(row=index, column=1, value=link)  # Assuming you want to save in the first column

# Save the workbook
exl.save('ids.xlsx')

# Закрыть браузер
input("Формирование ссылок завершено. Нажмите Enter для закрытия окна")
driver.quit()

# softy_plug